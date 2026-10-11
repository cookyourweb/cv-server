"""Read an uploaded CV (PDF or DOCX) into plain text, defensively (spec section 3).

The file is untrusted. The checks run in this order and stop at the first failure:
size (before any parsing), real type by magic bytes (never the name or Content-Type),
format limits (PDF pages, DOCX uncompressed size and entry count), text extraction,
minimum text. Failures raise `ErrorDeLectura` with a fixed `codigo`; the exception
carries no message, no content and no chained cause, so nothing the document says can
reach a log (REQ-3.5, see privacidad.py).

Codes and the HTTP status the route gives them:
    tamano 413 · tipo 415 · ilegible 422 · largo 422 · sin_texto 422 (use the manual form)
"""
import io
import re
import warnings
import zipfile
from dataclasses import dataclass

import privacidad  # noqa: F401  (importing it silences the parsing libraries' loggers)

LIMITE_BYTES = 2 * 1024 * 1024          # raw upload, checked before parsing (REQ-3.1)
MAX_PAGINAS = 10                        # PDF
MAX_DESCOMPRIMIDO = 10 * 1024 * 1024    # DOCX, sum of the uncompressed entries (zip bomb guard)
MAX_ENTRADAS_ZIP = 500                  # DOCX, number of entries
MIN_CARACTERES = 300                    # below this the file is treated as a scan (REQ-3.3)
MAX_CARACTERES = 200_000                # a CV is a few thousand; this only stops abuse

_PDF = b"%PDF-"
_ZIP = b"PK\x03\x04"
_DOCX_PRINCIPAL = "word/document.xml"


class ErrorDeLectura(Exception):
    """The CV cannot be used. `codigo` says why; the message is fixed and generic."""

    def __init__(self, codigo: str):
        super().__init__(codigo)
        self.codigo = codigo


@dataclass(frozen=True)
class CVLeido:
    formato: str    # "pdf" | "docx"
    texto: str
    idioma: str     # "es" | "en"

    @property
    def caracteres(self) -> int:
        return len(self.texto)


# --- language -------------------------------------------------------------------------------
# server.detectar_idioma is tuned for job offers (offer vocabulary, many ties) and lives in
# server.py, which imports this package's routes, so it cannot be imported from here. A CV is
# a different text, so this counts very common function words of each language instead.
# Deterministic; a tie (or no signal) is Spanish, the same default the offers use.
_PALABRAS_ES = frozenset("el la los las de del y en con por para un una que se su al como más".split())
_PALABRAS_EN = frozenset("the and of in with for to a an that is at as on by from my our".split())
_ACENTOS_ES = re.compile(r"[ñáéíóúü¿¡]", re.IGNORECASE)
_PALABRA = re.compile(r"[a-záéíóúñü]+")


def detectar_idioma(texto: str) -> str:
    palabras = _PALABRA.findall((texto or "").lower())
    es = sum(1 for p in palabras if p in _PALABRAS_ES) + len(_ACENTOS_ES.findall(texto or ""))
    en = sum(1 for p in palabras if p in _PALABRAS_EN)
    return "en" if en > es else "es"


# --- extraction -----------------------------------------------------------------------------------

def _extraer_pdf(datos: bytes) -> str:
    from pypdf import PdfReader
    lector = PdfReader(io.BytesIO(datos), strict=False)
    if lector.is_encrypted:
        raise ErrorDeLectura("ilegible")
    if len(lector.pages) > MAX_PAGINAS:
        raise ErrorDeLectura("largo")
    return "\n".join((pagina.extract_text() or "") for pagina in lector.pages)


def _extraer_docx(datos: bytes) -> str:
    from docx import Document
    documento = Document(io.BytesIO(datos))
    lineas = [p.text for p in documento.paragraphs]
    for tabla in documento.tables:
        for fila in tabla.rows:
            lineas.extend(celda.text for celda in fila.cells)
    return "\n".join(lineas)


def _inspeccionar_zip(datos: bytes) -> bool:
    """True for a DOCX within the limits, False for another kind of zip, None for a
    broken archive or one that breaks the limits (zip bomb, too many entries)."""
    try:
        with zipfile.ZipFile(io.BytesIO(datos)) as z:
            entradas = z.infolist()
            nombres = {e.filename for e in entradas}
            if len(entradas) > MAX_ENTRADAS_ZIP or sum(e.file_size for e in entradas) > MAX_DESCOMPRIMIDO:
                return None
            return _DOCX_PRINCIPAL in nombres
    except Exception:
        return None


def _limpiar(texto: str) -> str:
    texto = texto.replace("\x00", "")
    texto = "\n".join(linea.rstrip() for linea in texto.splitlines())
    return re.sub(r"\n{3,}", "\n\n", texto).strip()


def leer_cv(datos: bytes) -> CVLeido:
    """Validate and read an uploaded CV. Raises ErrorDeLectura; never anything else."""
    if len(datos) > LIMITE_BYTES:
        raise ErrorDeLectura("tamano")

    if datos.startswith(_PDF):
        formato, extraer = "pdf", _extraer_pdf
    elif datos.startswith(_ZIP):
        veredicto = _inspeccionar_zip(datos)
        if veredicto is None:
            raise ErrorDeLectura("ilegible")
        if not veredicto:
            raise ErrorDeLectura("tipo")
        formato, extraer = "docx", _extraer_docx
    else:
        raise ErrorDeLectura("tipo")

    # The parser's exception is caught and swallowed, and ours is raised AFTER the except
    # block: raising inside it would chain the parser's message to ours.
    codigo, crudo = None, ""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            crudo = extraer(datos)
        except ErrorDeLectura as error:
            codigo = error.codigo
        except Exception:
            codigo = "ilegible"
    if codigo:
        raise ErrorDeLectura(codigo)

    texto = _limpiar(crudo)
    if len(texto) > MAX_CARACTERES:
        raise ErrorDeLectura("largo")
    if len(texto) < MIN_CARACTERES:
        raise ErrorDeLectura("sin_texto")
    return CVLeido(formato, texto, detectar_idioma(texto))
